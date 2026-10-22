from datetime import timedelta
import random
from io import BytesIO

from django.core.cache import cache
from django.core.exceptions import ObjectDoesNotExist
from django.db import transaction
from django.db.models import Q
from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from rest_framework import viewsets, permissions, status
from rest_framework.decorators import action
from rest_framework.response import Response
from rest_framework.throttling import AnonRateThrottle
from rest_framework.views import APIView

from .models import Center, Branch, Subject, Level, Student, Question, Result, StudentAnswer, MentalTask
from .tests_loader import SUBJECT_ALIASES, name_key, normalize_level_key

def is_mental_subject(subject_name):
    normalized = subject_name.lower().replace("'", '').replace('‘', '').replace('’', '')
    return 'mental' in normalized


def get_exam_duration_minutes(student):
    if is_mental_subject(student.subject.name):
        return 5
    try:
        return int(student.level.duration_minutes or 60)
    except Exception:
        return 60


def get_exam_duration_seconds(student):
    return max(1, int(get_exam_duration_minutes(student) * 60))


def get_test_versions(student):
    versions = list(
        Question.objects.filter(subject=student.subject, level=student.level)
        .values_list('version', flat=True)
        .distinct()
        .order_by('version')
    )
    return [int(version or 1) for version in versions]


def exam_questions_cache_key(subject_id, level_id, version):
    return f'exam-questions:{subject_id}:{level_id}:{version}'


def get_exam_questions_data(student, version):
    """Imtihon savollari (to'g'ri javobsiz) qisqa muddatli keshdan olinadi.

    Yuzlab o'quvchi bir vaqtda kirganda har biri uchun bazadan bir xil savollarni
    qayta o'qimaslik uchun. Savol o'zgartirilsa kesh tozalanadi (signals.py),
    boshqa jarayonlarda esa ko'pi bilan 60 soniyada yangilanadi.
    """
    key = exam_questions_cache_key(student.subject_id, student.level_id, version)
    data = cache.get(key)
    if data is None:
        questions = Question.objects.filter(
            subject_id=student.subject_id,
            level_id=student.level_id,
            version=version,
        ).order_by('id')
        data = [dict(item) for item in QuestionForExamSerializer(questions, many=True).data]
        cache.set(key, data, 60)
    return data


def build_version_choices(student):
    choices = []
    for version in get_test_versions(student):
        count = len(get_exam_questions_data(student, version))
        if count:
            choices.append({
                'value': version,
                'label': 'Test',
                'questions_count': count,
            })
    return choices


def normalize_requested_version(raw_version):
    try:
        version = int(raw_version)
    except (TypeError, ValueError):
        return None
    return version if version > 0 else None


def normalize_progress_answers(raw_answers, is_mental=False):
    """Clean progress answers before saving them on the student row.

    Normal tests are saved as {question_id: 'A'}. Mental answers are saved as
    {task_id: '123'} so they can be restored even from another browser/device.
    """
    if not isinstance(raw_answers, dict):
        return {}

    cleaned = {}
    for key, value in raw_answers.items():
        key_text = str(key or '').strip()
        if not key_text:
            continue

        if is_mental:
            answer_text = str(value if value is not None else '').strip()
            if answer_text:
                cleaned[key_text] = answer_text[:30]
        else:
            answer_text = str(value if value is not None else '').strip().upper()
            if answer_text in ['A', 'B', 'C', 'D']:
                cleaned[key_text] = answer_text
    return cleaned


def get_resume_remaining_seconds(student):
    duration_seconds = get_exam_duration_seconds(student)
    saved_remaining = student.progress_remaining_seconds
    if saved_remaining is not None:
        try:
            return max(0, min(duration_seconds, int(saved_remaining)))
        except (TypeError, ValueError):
            pass

    if student.started_at:
        elapsed = int((timezone.now() - student.started_at).total_seconds())
        return max(0, duration_seconds - max(0, elapsed))

    return duration_seconds




def apply_visible_score_penalty(correct_count, total_questions):
    """Official 1-point penalty used for all visible saved results.

    The stored result is the final visible score shown to students and admins.
    It never goes below 0 and the percent is calculated from this final score.
    """
    total = max(0, int(total_questions or 0))
    raw = max(0, int(correct_count or 0))
    final_score = max(raw - 1, 0) if total else 0
    percent = (final_score / total * 100) if total else 0
    return final_score, round(percent, 2)

def calculate_spent_seconds(student, finished_at, submitted_remaining_seconds=None):
    """Return real working time shown in admin.

    If the student leaves the page and later enters the same code again, the
    frontend/server preserve remaining_seconds. So admin time should be based
    on the preserved timer, not just wall-clock time between started/finished.
    """
    duration_seconds = get_exam_duration_seconds(student)
    remaining = submitted_remaining_seconds

    if remaining is None:
        remaining = student.progress_remaining_seconds

    try:
        remaining = int(remaining)
        remaining = max(0, min(duration_seconds, remaining))
        return max(0, duration_seconds - remaining)
    except (TypeError, ValueError):
        pass

    started_at = student.started_at or finished_at
    return max(0, min(duration_seconds, int((finished_at - started_at).total_seconds())))


MENTAL_ARITHMETIC_SEQUENCES = [
    [22, -11, -11, 55, 44],
    [33, 66, -99, 77, 11],
    [66, 33, -77, -11, 22],
    [33, 66, -55, 55, -11],
    [11, 66, -11, 33, -33],
    [99, -33, -55, -11, 44],
    [22, 77, -44, -55, 55],
    [55, 44, -22, -66, 66],
    [22, 22, 55, -77, 11],
    [22, 55, 22, -88, -11],
    [33, 11, 55, -11, 11],
    [77, 11, -22, 33, -66],
    [55, -55, 55, 11, 33],
    [99, -22, 22, -99, 77],
    [33, -33, 22, -22, 33],
    [66, -11, -55, 33, -33],
    [77, -55, -22, 77, 22],
    [99, -22, 22, -66, 11],
    [66, 22, 11, -22, 11],
    [11, -11, 55, -55, 99],
    [66, 33, -66, -11, 55],
    [66, -11, -55, 77, 11],
    [44, 55, -44, -55, 99],
    [33, -11, 11, 55, -88],
    [33, 55, 11, -77, 66],
    [88, -33, -55, 55, 44],
    [11, -11, 88, 11, -44],
    [33, 66, -88, 33, 55],
    [22, 66, -33, -55, 44],
    [22, 77, -88, -11, 55],
    [77, -55, 11, 66, -11],
    [55, -55, 33, 11, -44],
    [66, -55, -11, 44, 55],
    [55, -55, 88, -22, -11],
    [11, -11, 88, -88, 44],
    [66, 11, 22, -55, -22],
    [44, 55, -44, -55, 99],
    [44, -33, -11, 22, -22],
    [88, -33, -55, 22, -22],
    [55, -55, 99, -88, -11],
    [55, -55, 99, -33, 11],
    [66, 33, -77, 55, 11],
    [33, 55, -11, -22, -55],
    [99, -33, 11, 22, -44],
    [11, 33, -33, 66, -22],
    [33, 11, 55, -77, 66],
    [55, 33, 11, -22, 22],
    [33, 11, 55, -33, -55],
    [99, -11, -77, -11, 33],
    [77, 22, -22, -55, 55],
    [55, -55, 44, -33, 66],
    [99, -88, -11, 99, -44],
    [55, 11, 33, -66, 66],
    [77, 22, -11, 11, -88],
    [44, -11, 55, -88, 33],
    [66, 33, -11, 11, -33],
    [22, 77, -99, 11, -11],
    [55, -55, 66, -11, 11],
    [44, -44, 11, -11, 99],
    [22, 55, 11, 11, -33],
    [66, -55, -11, 44, -33],
    [11, 33, -33, 88, -22],
    [44, 55, -66, -11, -22],
    [77, 22, -22, 11, 11],
    [66, -11, 22, 11, 11],
    [44, 55, -11, -22, 11],
    [22, 77, -33, -66, 33],
    [33, -22, -11, 44, 55],
    [11, 88, -33, 33, -99],
    [55, 33, -88, 55, -55],
    [77, 22, -66, 66, -11],
    [22, 22, 55, -33, 11],
    [99, -33, 22, 11, -66],
    [33, -33, 99, -11, -77],
    [44, -33, -11, 66, -66],
    [66, 11, -77, 88, 11],
    [88, 11, -99, 33, -33],
    [44, -11, 66, -77, 77],
    [77, 11, 11, -88, 88],
    [77, 22, -44, -55, 99],
]

def format_mental_number(value, is_first=False):
    return str(value)


def generate_mental_tasks(student, count=None):
    existing_tasks = list(MentalTask.objects.filter(student=student).order_by('task_order'))
    if existing_tasks:
        return existing_tasks

    created = []
    sequences = MENTAL_ARITHMETIC_SEQUENCES if count is None else MENTAL_ARITHMETIC_SEQUENCES[:count]

    for order, sequence in enumerate(sequences, start=1):
        flashes = [format_mental_number(value, index == 0) for index, value in enumerate(sequence)]
        correct_answer = sum(sequence)
        task = MentalTask.objects.create(
            student=student,
            task_order=order,
            flashes=flashes,
            expression=' '.join(flashes),
            correct_answer=correct_answer,
        )
        created.append(task)
    return created





from .serializers import (
    CenterSerializer,
    BranchSerializer,
    SubjectSerializer,
    LevelSerializer,
    StudentSerializer,
    QuestionAdminSerializer,
    QuestionForExamSerializer,
    ResultSerializer,
)




DEFAULT_SUBJECT_LEVELS_FOR_FORMS = {
    'Koreys tili': ['Koreys tili 1', 'Koreys tili 2'],
}


def resolve_subject_and_level(subject_name, level_name):
    """Excel'dagi fan/daraja nomini bazadagi testli fan/darajaga moslaydi.

    Masalan: "INGLIZ TILI" -> "English", "MATEM" + "3-SINF" -> "Matematika" / "3-sinf".
    Mos keladigan nom topilmasa, avvalgidek yangi fan/daraja yaratiladi.
    """
    subject_key = name_key(subject_name)
    canonical = SUBJECT_ALIASES.get(subject_key)
    subject = None
    for item in Subject.objects.all():
        if name_key(item.name) == name_key(canonical or subject_name):
            subject = item
            break
    if subject is None:
        subject, _ = Subject.objects.get_or_create(name=canonical or subject_name)

    level_key = normalize_level_key(level_name)
    for item in Level.objects.filter(subject=subject):
        if name_key(item.name) == level_key:
            return subject, item
    level, _ = Level.objects.get_or_create(subject=subject, name=level_name, defaults={'duration_minutes': 30})
    return subject, level


def ensure_subject_levels_for_forms():
    """Keep admin create-student dropdowns from missing newly added subjects.

    Railway runs migrations before the app starts, but older deployments/databases can
    still miss a subject until seed_demo finishes. This lightweight guard guarantees
    that /subjects/ and /levels/ can show Koreys tili immediately.
    """
    for subject_name, level_names in DEFAULT_SUBJECT_LEVELS_FOR_FORMS.items():
        subject, _ = Subject.objects.get_or_create(name=subject_name)
        for level_name in level_names:
            Level.objects.get_or_create(
                subject=subject,
                name=level_name,
                defaults={'duration_minutes': 30},
            )


class IsAdminUser(permissions.BasePermission):
    def has_permission(self, request, view):
        return bool(request.user and request.user.is_authenticated and request.user.is_staff)


def get_user_profile(user):
    return getattr(user, 'admin_profile', None) if user and user.is_authenticated else None


def get_user_branch(user):
    return getattr(get_user_profile(user), 'branch', '')


def get_user_center(user):
    profile = get_user_profile(user)
    return getattr(profile, 'center', None)


def is_main_admin(user):
    profile = get_user_profile(user)
    return bool(user and user.is_authenticated and user.is_staff and (user.is_superuser or not profile or (not profile.center_id and not profile.branch)))


def can_manage_students(user):
    return bool(is_main_admin(user) or get_user_center(user))


def branch_key(text):
    return str(text or '').strip().lower().replace("'", '').replace('‘', '').replace('’', '').replace('`', '')


def normalize_branch(value):
    clean = str(value or '').strip()
    if not clean:
        return ''

    try:
        existing = Branch.objects.filter(name__iexact=clean).first()
        if existing:
            return existing.name
    except Exception:
        pass
    return clean


def ensure_branch_exists(branch_name):
    clean = normalize_branch(branch_name)
    if not clean:
        return ''
    Branch.objects.get_or_create(name=clean)
    return clean


class CenterViewSet(viewsets.ModelViewSet):
    queryset = Center.objects.all()
    serializer_class = CenterSerializer
    permission_classes = [IsAdminUser]

    def get_queryset(self):
        qs = super().get_queryset()
        user_center = get_user_center(self.request.user)
        if not is_main_admin(self.request.user) and user_center:
            qs = qs.filter(id=user_center.id)
        elif not is_main_admin(self.request.user):
            qs = qs.none()
        return qs.order_by('name')

    def create(self, request, *args, **kwargs):
        if not is_main_admin(request.user):
            return Response({'detail': 'Sizga o‘quv markaz qo‘shishga ruxsat yo‘q.'}, status=status.HTTP_403_FORBIDDEN)
        name = str(request.data.get('name') or '').strip()
        if not name:
            return Response({'name': 'O‘quv markaz nomini kiriting.'}, status=status.HTTP_400_BAD_REQUEST)
        existing = Center.objects.filter(name__iexact=name).first()
        if existing:
            serializer = self.get_serializer(existing)
            return Response(serializer.data, status=status.HTTP_200_OK)
        return super().create(request, *args, **kwargs)

    def update(self, request, *args, **kwargs):
        if not is_main_admin(request.user):
            return Response({'detail': 'Sizga o‘quv markazni tahrirlashga ruxsat yo‘q.'}, status=status.HTTP_403_FORBIDDEN)
        return super().update(request, *args, **kwargs)

    def partial_update(self, request, *args, **kwargs):
        if not is_main_admin(request.user):
            return Response({'detail': 'Sizga o‘quv markazni tahrirlashga ruxsat yo‘q.'}, status=status.HTTP_403_FORBIDDEN)
        return super().partial_update(request, *args, **kwargs)

    def destroy(self, request, *args, **kwargs):
        if not is_main_admin(request.user):
            return Response({'detail': 'Sizga o‘quv markazni o‘chirishga ruxsat yo‘q.'}, status=status.HTTP_403_FORBIDDEN)
        center = self.get_object()
        if Student.objects.filter(center=center).exists():
            return Response({'detail': 'Bu o‘quv markazda o‘quvchilar bor. Avval ularni boshqa markazga o‘tkazing.'}, status=status.HTTP_400_BAD_REQUEST)
        return super().destroy(request, *args, **kwargs)

    @action(detail=False, methods=['get'], url_path='summary')
    def summary(self, request):
        data = []
        for center in self.get_queryset():
            students_qs = Student.objects.filter(center=center)
            results_qs = Result.objects.filter(student__center=center)
            results = list(results_qs.values_list('percent', flat=True))
            avg_percent = round(sum(results) / len(results), 1) if results else 0
            data.append({
                'id': center.id,
                'name': center.name,
                'students_count': students_qs.count(),
                'not_started_count': students_qs.filter(status=Student.Status.NOT_STARTED).count(),
                'in_progress_count': students_qs.filter(status=Student.Status.IN_PROGRESS).count(),
                'completed_count': students_qs.filter(status=Student.Status.COMPLETED).count(),
                'results_count': len(results),
                'average_percent': avg_percent,
            })
        return Response(data)


class BranchViewSet(viewsets.ModelViewSet):
    queryset = Branch.objects.all()
    serializer_class = BranchSerializer
    permission_classes = [IsAdminUser]

    def create(self, request, *args, **kwargs):
        if not is_main_admin(request.user):
            return Response({'detail': 'Sizga filial qo‘shishga ruxsat yo‘q.'}, status=status.HTTP_403_FORBIDDEN)
        name = str(request.data.get('name') or '').strip()
        if not name:
            return Response({'name': 'Filial nomini kiriting.'}, status=status.HTTP_400_BAD_REQUEST)
        existing = Branch.objects.filter(name__iexact=name).first()
        if existing:
            serializer = self.get_serializer(existing)
            return Response(serializer.data, status=status.HTTP_200_OK)
        return super().create(request, *args, **kwargs)

    def update(self, request, *args, **kwargs):
        if not is_main_admin(request.user):
            return Response({'detail': 'Sizga filialni tahrirlashga ruxsat yo‘q.'}, status=status.HTTP_403_FORBIDDEN)
        return super().update(request, *args, **kwargs)

    def partial_update(self, request, *args, **kwargs):
        if not is_main_admin(request.user):
            return Response({'detail': 'Sizga filialni tahrirlashga ruxsat yo‘q.'}, status=status.HTTP_403_FORBIDDEN)
        return super().partial_update(request, *args, **kwargs)

    def destroy(self, request, *args, **kwargs):
        if not is_main_admin(request.user):
            return Response({'detail': 'Sizga filialni o‘chirishga ruxsat yo‘q.'}, status=status.HTTP_403_FORBIDDEN)
        branch = self.get_object()
        if Student.objects.filter(branch=branch.name).exists():
            return Response({'detail': 'Bu filialda o‘quvchilar bor. Avval ularni boshqa filialga o‘tkazing.'}, status=status.HTTP_400_BAD_REQUEST)
        return super().destroy(request, *args, **kwargs)


class SubjectViewSet(viewsets.ModelViewSet):
    queryset = Subject.objects.all()
    serializer_class = SubjectSerializer
    permission_classes = [IsAdminUser]

    def get_queryset(self):
        ensure_subject_levels_for_forms()
        return Subject.objects.all().order_by('name')


class LevelViewSet(viewsets.ModelViewSet):
    queryset = Level.objects.select_related('subject').all()
    serializer_class = LevelSerializer
    permission_classes = [IsAdminUser]

    def get_queryset(self):
        ensure_subject_levels_for_forms()
        qs = Level.objects.select_related('subject').all()
        subject_id = self.request.query_params.get('subject')
        if subject_id:
            qs = qs.filter(subject_id=subject_id)
        return qs


class StudentViewSet(viewsets.ModelViewSet):
    # 'result' ham birga olinadi: ro'yxatda har bir o'quvchi uchun alohida so'rov ketmasin.
    queryset = Student.objects.select_related('subject', 'level', 'center', 'result').all()
    serializer_class = StudentSerializer
    permission_classes = [IsAdminUser]

    def get_queryset(self):
        qs = super().get_queryset()
        q = self.request.query_params.get('q')
        status_value = self.request.query_params.get('status')
        subject = self.request.query_params.get('subject')
        level = self.request.query_params.get('level')
        center = self.request.query_params.get('center')
        branch = normalize_branch(self.request.query_params.get('branch'))
        user_center = get_user_center(self.request.user)
        user_branch = get_user_branch(self.request.user)

        if not is_main_admin(self.request.user) and user_center:
            qs = qs.filter(center=user_center)
        elif not is_main_admin(self.request.user) and user_branch:
            qs = qs.filter(branch=user_branch)
        elif center:
            qs = qs.filter(center_id=center)

        if branch and is_main_admin(self.request.user):
            qs = qs.filter(branch=branch)

        if q:
            qs = qs.filter(Q(first_name__icontains=q) | Q(last_name__icontains=q) | Q(code__icontains=q))
        if status_value:
            qs = qs.filter(status=status_value)
        if subject:
            qs = qs.filter(subject_id=subject)
        if level:
            qs = qs.filter(level_id=level)
        return qs.order_by('-created_at')

    def create(self, request, *args, **kwargs):
        if not can_manage_students(request.user):
            return Response({'detail': 'Sizga o‘quvchi yaratishga ruxsat yo‘q.'}, status=status.HTTP_403_FORBIDDEN)
        return super().create(request, *args, **kwargs)

    def perform_create(self, serializer):
        user_center = get_user_center(self.request.user)
        branch = ensure_branch_exists(serializer.validated_data.get('branch')) or ensure_branch_exists(get_user_branch(self.request.user)) or 'Boshqa'
        save_kwargs = {
            'branch': branch,
            'status': Student.Status.NOT_STARTED,
            'started_at': None,
            'finished_at': None,
            'is_used': False,
            'selected_version': None,
            'progress_answers': {},
            'progress_remaining_seconds': None,
            'progress_current_index': 0,
            'progress_updated_at': None,
        }
        if not is_main_admin(self.request.user) and user_center:
            save_kwargs['center'] = user_center
        serializer.save(**save_kwargs)

    def update(self, request, *args, **kwargs):
        if not can_manage_students(request.user):
            return Response({'detail': 'Sizga o‘quvchini tahrirlashga ruxsat yo‘q.'}, status=status.HTTP_403_FORBIDDEN)
        return super().update(request, *args, **kwargs)

    def partial_update(self, request, *args, **kwargs):
        if not can_manage_students(request.user):
            return Response({'detail': 'Sizga o‘quvchini tahrirlashga ruxsat yo‘q.'}, status=status.HTTP_403_FORBIDDEN)
        return super().partial_update(request, *args, **kwargs)

    def destroy(self, request, *args, **kwargs):
        if not can_manage_students(request.user):
            return Response({'detail': 'Sizga o‘quvchini o‘chirishga ruxsat yo‘q.'}, status=status.HTTP_403_FORBIDDEN)
        return super().destroy(request, *args, **kwargs)

    @action(detail=False, methods=['post'], url_path='bulk-delete')
    def bulk_delete(self, request):
        if not can_manage_students(request.user):
            return Response({'detail': 'Sizga o‘quvchilarni o‘chirishga ruxsat yo‘q.'}, status=status.HTTP_403_FORBIDDEN)

        raw_ids = request.data.get('ids', [])
        if not isinstance(raw_ids, list):
            return Response({'detail': 'ids ro‘yxat bo‘lishi kerak.'}, status=status.HTTP_400_BAD_REQUEST)

        ids = []
        for item in raw_ids:
            try:
                ids.append(int(item))
            except (TypeError, ValueError):
                continue

        if not ids:
            return Response({'detail': 'O‘chirish uchun o‘quvchi tanlanmadi.'}, status=status.HTTP_400_BAD_REQUEST)

        qs = self.get_queryset().filter(id__in=ids)
        deleted_count = qs.count()
        qs.delete()
        return Response({'deleted_count': deleted_count})

    @action(detail=False, methods=['post'], url_path='import-excel')
    def import_excel(self, request):
        if not can_manage_students(request.user):
            return Response({'detail': 'Sizga Excel orqali o‘quvchi yaratishga ruxsat yo‘q.'}, status=status.HTTP_403_FORBIDDEN)
        file_obj = request.FILES.get('file')
        if not file_obj:
            return Response({'detail': 'Excel file yuborilmadi.'}, status=status.HTTP_400_BAD_REQUEST)

        try:
            wb = load_workbook(file_obj)
            ws = wb.active
        except Exception:
            return Response({'detail': 'Excel faylni o‘qib bo‘lmadi.'}, status=status.HTTP_400_BAD_REQUEST)

        def normalize_header(value):
            text = str(value or '').strip().lower()
            replacements = {
                '‘': "'",
                '’': "'",
                '`': "'",
                'ʼ': "'",
                '№': 'no',
                '#': 'no',
            }
            for old, new_value in replacements.items():
                text = text.replace(old, new_value)
            text = ' '.join(text.split())
            return text

        headers = [normalize_header(cell.value) for cell in ws[1]]
        header_map = {name: idx for idx, name in enumerate(headers) if name}

        def cell_value(row, names):
            for name in names:
                idx = header_map.get(normalize_header(name))
                if idx is not None and idx < len(row):
                    value = row[idx]
                    return str(value).strip() if value is not None else ''
            return ''

        def split_full_name(full_name):
            parts = str(full_name or '').strip().split()
            if not parts:
                return '', ''
            if len(parts) == 1:
                return parts[0], ''
            return parts[0], ' '.join(parts[1:])

        created = []
        errors = []
        user_center = get_user_center(request.user)
        user_branch = get_user_branch(request.user)

        with transaction.atomic():
            for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                full_name = cell_value(row, [
                    'Ism familya',
                    'Ism familyasi',
                    'Ism familiya',
                    'Ism familiyasi',
                    'Familiya ism',
                    'Ism Familyasi',
                    'F.I.Sh',
                    'FISH',
                    'FIO',
                    'Oquvchi',
                    "O'quvchi",
                    'O‘quvchi',
                ])
                first_name = cell_value(row, ['Ism', 'first_name', 'First name'])
                last_name = cell_value(row, ['Familya', 'last_name', 'Last name'])

                if full_name and not first_name and not last_name:
                    first_name, last_name = split_full_name(full_name)

                subject_name = cell_value(row, ['Fan', 'Subject'])
                level_name = cell_value(row, ['Daraja', 'Level'])
                center_name = cell_value(row, ["O'quv markaz", 'O‘quv markaz', 'Oquv markaz', "O'quv markazi", 'O‘quv markazi', 'Oquv markazi', 'Center'])
                branch_name = normalize_branch(cell_value(row, ['Filial', 'Branch']))

                if not is_main_admin(request.user) and user_center:
                    center = user_center
                    center_name = user_center.name
                    if not branch_name:
                        branch_name = user_branch
                else:
                    center = None

                if not any([full_name, first_name, last_name, subject_name, level_name, center_name, branch_name]):
                    continue

                if not all([first_name, subject_name, level_name, center_name, branch_name]):
                    errors.append({
                        'row': row_num,
                        'error': "Majburiy ustunlar: Ism familya, Fan, Daraja, O'quv markaz, Filial.",
                    })
                    continue

                branch_name = ensure_branch_exists(branch_name)

                subject, level = resolve_subject_and_level(subject_name, level_name)
                if center is None:
                    center, _ = Center.objects.get_or_create(name=center_name)

                student = Student.objects.create(
                    first_name=first_name,
                    last_name=last_name,
                    subject=subject,
                    level=level,
                    center=center,
                    branch=branch_name,
                    status=Student.Status.NOT_STARTED,
                    started_at=None,
                    finished_at=None,
                    is_used=False,
                    selected_version=None,
                    progress_answers={},
                    progress_remaining_seconds=None,
                    progress_current_index=0,
                    progress_updated_at=None,
                )
                created.append(student)

        return Response({
            'created_count': len(created),
            'errors': errors,
            'students': StudentSerializer(created, many=True).data,
        }, status=status.HTTP_201_CREATED)

    @action(detail=False, methods=['get'], url_path='export-excel')
    def export_excel(self, request):
        students = self.get_queryset().select_related('subject', 'level', 'center', 'result')

        wb = Workbook()
        ws = wb.active
        ws.title = 'Barcha oquvchilar'

        headers = ['№', 'Ism familyasi', 'Fani', 'Darajasi', 'Yoshi', "O'quv markazi", 'Filial', 'Code', 'Status', 'Natijasi', 'Sarflagan vaqt']
        ws.append(headers)

        header_fill = PatternFill('solid', fgColor='1F4E79')
        for cell in ws[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center')

        status_labels = {
            Student.Status.NOT_STARTED: 'Ishlamagan',
            Student.Status.IN_PROGRESS: 'Ishlayapti',
            Student.Status.COMPLETED: 'Ishlab bo‘ldi',
        }

        for idx, student in enumerate(students, start=1):
            try:
                result = student.result
            except ObjectDoesNotExist:
                result = None

            natija = '—'
            spent_time = '—'
            if result:
                natija = f'{result.correct_count}/{result.total_questions} ta / {result.percent:.1f}%'
                spent_time = str(timedelta(seconds=result.spent_seconds))

            ws.append([
                idx,
                student.full_name,
                student.subject.name,
                student.level.name,
                student.age or '—',
                student.center.name,
                student.branch,
                student.code,
                status_labels.get(student.status, student.status),
                natija,
                spent_time,
            ])

        ws.freeze_panes = 'A2'
        for column_cells in ws.columns:
            max_length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
            ws.column_dimensions[get_column_letter(column_cells[0].column)].width = min(max_length + 4, 42)

        output = BytesIO()
        wb.save(output)
        output.seek(0)

        response = HttpResponse(
            output.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="barcha_oquvchilar.xlsx"'
        return response

class QuestionViewSet(viewsets.ModelViewSet):
    queryset = Question.objects.select_related('subject', 'level').all()
    serializer_class = QuestionAdminSerializer
    permission_classes = [IsAdminUser]

    def get_queryset(self):
        qs = super().get_queryset()
        subject = self.request.query_params.get('subject')
        level = self.request.query_params.get('level')
        if subject:
            qs = qs.filter(subject_id=subject)
        if level:
            qs = qs.filter(level_id=level)
        return qs.order_by('subject__name', 'level__name', 'id')

    def create(self, request, *args, **kwargs):
        if not is_main_admin(request.user):
            return Response({'detail': 'Sizga test qo‘shishga ruxsat yo‘q.'}, status=status.HTTP_403_FORBIDDEN)
        return super().create(request, *args, **kwargs)

    def update(self, request, *args, **kwargs):
        if not is_main_admin(request.user):
            return Response({'detail': 'Sizga testni tahrirlashga ruxsat yo‘q.'}, status=status.HTTP_403_FORBIDDEN)
        return super().update(request, *args, **kwargs)

    def partial_update(self, request, *args, **kwargs):
        if not is_main_admin(request.user):
            return Response({'detail': 'Sizga testni tahrirlashga ruxsat yo‘q.'}, status=status.HTTP_403_FORBIDDEN)
        return super().partial_update(request, *args, **kwargs)

    def destroy(self, request, *args, **kwargs):
        if not is_main_admin(request.user):
            return Response({'detail': 'Sizga testni o‘chirishga ruxsat yo‘q.'}, status=status.HTTP_403_FORBIDDEN)
        return super().destroy(request, *args, **kwargs)


class ResultViewSet(viewsets.ReadOnlyModelViewSet):
    queryset = Result.objects.select_related('student', 'student__subject', 'student__level', 'student__center').prefetch_related('answers__question', 'mental_answers').all()
    serializer_class = ResultSerializer
    permission_classes = [IsAdminUser]

    def get_queryset(self):
        qs = super().get_queryset()
        center = self.request.query_params.get('center')
        branch = normalize_branch(self.request.query_params.get('branch'))
        user_center = get_user_center(self.request.user)
        user_branch = get_user_branch(self.request.user)

        if not is_main_admin(self.request.user) and user_center:
            qs = qs.filter(student__center=user_center)
        elif not is_main_admin(self.request.user) and user_branch:
            qs = qs.filter(student__branch=user_branch)
        elif center:
            qs = qs.filter(student__center_id=center)

        if branch and is_main_admin(self.request.user):
            qs = qs.filter(student__branch=branch)
        return qs.order_by('-created_at')

    @action(detail=False, methods=['get'], url_path='export-excel')
    def export_excel(self, request):
        results = list(self.get_queryset())

        wb = Workbook()
        default_ws = wb.active
        wb.remove(default_ws)

        headers = [
            '№', 'Ism Familya', 'Fan', 'Daraja', 'Version', "O'quv markaz", 'Filial', 'Status code',
            "Nechta to'g'ri", 'Jami savollar', 'Foiz', 'Boshlangan vaqti',
            'Tugatgan vaqti', 'Sarflagan vaqt'
        ]
        header_fill = PatternFill('solid', fgColor='1F4E79')

        def safe_sheet_title(name, used_titles):
            clean = str(name or 'Filialsiz').strip() or 'Filialsiz'
            for char in ['\\', '/', '*', '?', ':', '[', ']']:
                clean = clean.replace(char, '-')
            clean = clean[:31] or 'Filialsiz'
            title = clean
            counter = 2
            while title in used_titles:
                suffix = f' {counter}'
                title = f'{clean[:31 - len(suffix)]}{suffix}'
                counter += 1
            used_titles.add(title)
            return title

        grouped = {}
        for result in results:
            branch_name = result.student.branch or 'Filialsiz'
            grouped.setdefault(branch_name, []).append(result)

        if not grouped:
            grouped['Natijalar'] = []

        used_titles = set()
        for branch_name in sorted(grouped.keys(), key=lambda item: str(item).lower()):
            ws = wb.create_sheet(title=safe_sheet_title(branch_name, used_titles))
            ws.append(headers)

            for cell in ws[1]:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal='center')

            for idx, result in enumerate(grouped[branch_name], start=1):
                spent = str(timedelta(seconds=result.spent_seconds))
                ws.append([
                    idx,
                    result.student.full_name,
                    result.student.subject.name,
                    result.student.level.name,
                    f"Version {result.student.selected_version}" if result.student.selected_version else '—',
                    result.student.center.name,
                    result.student.branch,
                    result.student.code,
                    result.correct_count,
                    result.total_questions,
                    round(float(result.percent or 0), 1),
                    timezone.localtime(result.started_at).strftime('%Y-%m-%d %H:%M:%S') if result.started_at else '',
                    timezone.localtime(result.finished_at).strftime('%Y-%m-%d %H:%M:%S') if result.finished_at else '',
                    spent,
                ])

            ws.freeze_panes = 'A2'
            for column_cells in ws.columns:
                max_length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
                ws.column_dimensions[get_column_letter(column_cells[0].column)].width = min(max_length + 3, 40)

        output = BytesIO()
        wb.save(output)
        output.seek(0)

        response = HttpResponse(
            output.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="olimpiada_natijalari_filiallar.xlsx"'
        return response


    @action(detail=False, methods=['get'], url_path='mental-answers')
    def mental_answers(self, request):
        qs = self.get_queryset().filter(mental_answers__isnull=False).distinct().prefetch_related('mental_answers')
        serializer = self.get_serializer(qs, many=True)
        return Response(serializer.data)

    @action(detail=False, methods=['get'], url_path='mental-answers-export')
    def mental_answers_export(self, request):
        qs = self.get_queryset().filter(mental_answers__isnull=False).distinct().prefetch_related('mental_answers')

        wb = Workbook()
        ws = wb.active
        ws.title = 'Mental javoblari'

        headers = [
            '№', 'Ism Familya', 'Fan', 'Daraja', 'Version', "O'quv markaz", 'Filial', 'Status code',
            'Umumiy natija', 'Foiz', 'Sarflagan vaqt', 'Misol №', 'Misol', "O'quvchi javobi",
            "To'g'ri javob", 'Holat'
        ]
        ws.append(headers)

        header_fill = PatternFill('solid', fgColor='1F4E79')
        for cell in ws[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center')

        row_index = 1
        for result in qs:
            tasks = list(result.mental_answers.all().order_by('task_order'))
            if not tasks:
                continue
            spent = str(timedelta(seconds=result.spent_seconds))
            for task in tasks:
                ws.append([
                    row_index,
                    result.student.full_name,
                    result.student.subject.name,
                    result.student.level.name,
                    f"Version {result.student.selected_version}" if result.student.selected_version else '—',
                    result.student.center.name,
                    result.student.branch,
                    result.student.code,
                    f'{result.correct_count}/{result.total_questions}',
                    f'{result.percent:.1f}%',
                    spent,
                    task.task_order,
                    task.expression,
                    task.student_answer if task.student_answer is not None else '',
                    task.correct_answer,
                    "To'g'ri" if task.is_correct else "Noto'g'ri",
                ])
            row_index += 1

        for column_cells in ws.columns:
            max_length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
            ws.column_dimensions[get_column_letter(column_cells[0].column)].width = min(max_length + 3, 45)

        output = BytesIO()
        wb.save(output)
        output.seek(0)

        response = HttpResponse(
            output.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="mental_javoblari.xlsx"'
        return response


class PublicResultLookupAPIView(APIView):
    permission_classes = [permissions.AllowAny]

    def post(self, request):
        code = str(request.data.get('code', '')).strip()
        if not code:
            return Response({'detail': 'Status code kiriting.'}, status=status.HTTP_400_BAD_REQUEST)

        try:
            student = Student.objects.select_related('subject', 'level', 'center').get(code=code)
        except Student.DoesNotExist:
            return Response({'detail': 'Bunday code topilmadi.'}, status=status.HTTP_404_NOT_FOUND)

        if student.status != Student.Status.COMPLETED:
            return Response({'detail': 'Bu o‘quvchi hali testni yakunlamagan.'}, status=status.HTTP_400_BAD_REQUEST)

        try:
            result = student.result
        except ObjectDoesNotExist:
            return Response({'detail': 'Bu code uchun natija topilmadi.'}, status=status.HTTP_404_NOT_FOUND)

        is_mental = is_mental_subject(student.subject.name)
        mental_answers = []
        if is_mental:
            mental_answers = [
                {
                    'id': task.id,
                    'task_order': task.task_order,
                    'expression': task.expression,
                    'correct_answer': task.correct_answer,
                    'student_answer': task.student_answer,
                    'is_correct': task.is_correct,
                }
                for task in result.mental_answers.all().order_by('task_order')
            ]

        return Response({
            'student_full_name': student.full_name,
            'student_code': student.code,
            'subject_name': student.subject.name,
            'level_name': student.level.name,
            'selected_version': student.selected_version,
            'center_name': student.center.name,
            'branch': student.branch,
            'total_questions': result.total_questions,
            'correct_count': result.correct_count,
            'percent': result.percent,
            'spent_seconds': result.spent_seconds,
            'finished_at': result.finished_at,
            'is_mental': is_mental,
            'mental_answers': mental_answers,
        })

# --------------------------------------------------------------------------
# O'quvchining o'zi ro'yxatdan o'tishi (code'siz kirish)
# --------------------------------------------------------------------------
LEVEL_DISPLAY_ORDER = [
    'starter', 'beginner', 'elementary', 'pre-intermediate', 'intermediate',
    'upper-intermediate', 'advanced', 'a1', 'a2', 'b1', 'b2', 'c1',
    'html css', 'javascript', 'vue', 'python',
]


def level_sort_key(level):
    key = normalize_level_key(level.name)
    number = key.split('-')[0]
    if key.endswith('-sinf') and number.isdigit():
        return (0, int(number), '')
    if key in LEVEL_DISPLAY_ORDER:
        return (1, LEVEL_DISPLAY_ORDER.index(key), '')
    return (2, 0, key)


def build_registration_catalog():
    """Ro'yxatdan o'tish oynasi uchun faqat testi bor fan va darajalar.

    Mental arifmetikada savollar avtomatik yaratiladi, shuning uchun uning darajalari ham chiqadi.
    """
    data = cache.get('registration-catalog')
    if data is not None:
        return data
    levels = Level.objects.select_related('subject').filter(questions__isnull=False).distinct()
    by_subject = {}
    for level in levels:
        by_subject.setdefault(level.subject, []).append(level)
    for level in Level.objects.select_related('subject').all():
        if is_mental_subject(level.subject.name):
            by_subject.setdefault(level.subject, [])
            if level not in by_subject[level.subject]:
                by_subject[level.subject].append(level)
    data = {
        'subjects': [
            {
                'id': subject.id,
                'name': subject.name,
                'levels': [
                    {'id': level.id, 'name': level.name}
                    for level in sorted(subject_levels, key=level_sort_key)
                ],
            }
            for subject, subject_levels in sorted(by_subject.items(), key=lambda item: item[0].name.lower())
            if subject_levels
        ],
        'centers': list(Center.objects.order_by('name').values_list('name', flat=True)),
    }
    cache.set('registration-catalog', data, 60)
    return data


class RegistrationThrottle(AnonRateThrottle):
    # Bitta sinf bitta IP orqali kirsa ham yetadigan, lekin botni to'xtatadigan chegara.
    rate = '60/min'


class ExamCatalogAPIView(APIView):
    permission_classes = [permissions.AllowAny]
    authentication_classes = []

    def get(self, request):
        return Response(build_registration_catalog())


class ExamRegisterAPIView(APIView):
    permission_classes = [permissions.AllowAny]
    authentication_classes = []
    throttle_classes = [RegistrationThrottle]

    @transaction.atomic
    def post(self, request):
        full_name = ' '.join(str(request.data.get('full_name', '')).split())
        center_name = ' '.join(str(request.data.get('center', '')).split())
        errors = {}

        parts = full_name.split(' ')
        if len(full_name) < 3 or len(parts) < 2:
            errors['full_name'] = 'Ism va familiyani to‘liq kiriting.'
        elif len(full_name) > 200:
            errors['full_name'] = 'Ism familiya juda uzun.'

        if len(center_name) < 2:
            errors['center'] = 'O‘quv markaz nomini kiriting.'
        elif len(center_name) > 120:
            errors['center'] = 'O‘quv markaz nomi juda uzun.'

        try:
            age = int(request.data.get('age'))
            if not 4 <= age <= 99:
                raise ValueError
        except (TypeError, ValueError):
            age = None
            errors['age'] = 'Yoshni to‘g‘ri kiriting (4–99).'

        try:
            level_id = int(request.data.get('level'))
        except (TypeError, ValueError):
            level_id = None
        allowed_level_ids = {
            level['id'] for subject in build_registration_catalog()['subjects'] for level in subject['levels']
        }
        level = None
        if level_id in allowed_level_ids:
            level = Level.objects.select_related('subject').filter(id=level_id).first()
        if level is None:
            errors['level'] = 'Fan va darajani tanlang.'

        if errors:
            first_error = next(iter(errors.values()))
            return Response({'detail': first_error, 'errors': errors}, status=status.HTTP_400_BAD_REQUEST)

        # Markaz nomi katta-kichik harf farqisiz mavjudiga bog'lanadi, bo'lmasa yangisi ochiladi.
        center = next(
            (item for item in Center.objects.all() if name_key(item.name) == name_key(center_name)),
            None,
        )
        if center is None:
            center = Center.objects.create(name=center_name)
            cache.delete('registration-catalog')

        student = Student.objects.create(
            first_name=parts[0],
            last_name=' '.join(parts[1:]),
            subject=level.subject,
            level=level,
            center=center,
            age=age,
            is_self_registered=True,
        )
        return Response({
            'code': student.code,
            'full_name': student.full_name,
            'subject_name': level.subject.name,
            'level_name': level.name,
        }, status=status.HTTP_201_CREATED)


class ExamStartAPIView(APIView):
    permission_classes = [permissions.AllowAny]

    @transaction.atomic
    def post(self, request):
        code = str(request.data.get('code', '')).strip()
        requested_version = normalize_requested_version(request.data.get('version'))
        if not code:
            return Response({'detail': 'Status code kiriting.'}, status=status.HTTP_400_BAD_REQUEST)

        try:
            student = Student.objects.select_related('subject', 'level', 'center').select_for_update().get(code=code)
        except Student.DoesNotExist:
            return Response({'detail': 'Bunday code topilmadi.'}, status=status.HTTP_404_NOT_FOUND)

        if student.status == Student.Status.COMPLETED or student.is_used:
            return Response({'detail': 'Bu code oldin ishlatilgan.'}, status=status.HTTP_400_BAD_REQUEST)

        is_mental_exam = is_mental_subject(student.subject.name)
        duration_minutes = get_exam_duration_minutes(student)
        duration_seconds = get_exam_duration_seconds(student)

        def build_version_select_payload():
            choices = build_version_choices(student)
            if not choices:
                return None
            return {
                'mode': 'version_select',
                'student': StudentSerializer(student).data,
                'duration_minutes': duration_minutes,
                'available_versions': choices,
            }

        def build_mental_payload(started_at, resume=False):
            mental_tasks = generate_mental_tasks(student)
            return {
                'mode': 'mental',
                'student': StudentSerializer(student).data,
                'duration_minutes': 5,
                'remaining_seconds': get_resume_remaining_seconds(student),
                'started_at': started_at.isoformat(),
                'server_now': timezone.now().isoformat(),
                'resume': bool(resume),
                'saved_answers': student.progress_answers or {},
                'progress_current_index': student.progress_current_index or 0,
                'mental_tasks': [
                    {
                        'id': task.id,
                        'task_order': task.task_order,
                        'flashes': task.flashes,
                        'task_display_ms': 3000,
                    }
                    for task in mental_tasks
                ],
            }

        def build_test_payload(started_at, resume=False):
            exam_version = student.selected_version or requested_version or 1
            questions = get_exam_questions_data(student, exam_version)
            if not questions:
                return None
            return {
                'mode': 'test',
                'student': StudentSerializer(student).data,
                'duration_minutes': duration_minutes,
                'remaining_seconds': get_resume_remaining_seconds(student),
                'started_at': started_at.isoformat(),
                'server_now': timezone.now().isoformat(),
                'resume': bool(resume),
                'selected_version': exam_version,
                'saved_answers': student.progress_answers or {},
                'progress_current_index': student.progress_current_index or 0,
                'questions': questions,
            }

        if not is_mental_exam and student.status == Student.Status.NOT_STARTED:
            version_choices = build_version_choices(student)
            if not version_choices:
                return Response({'detail': 'Bu fan va daraja uchun testlar hali qo‘shilmagan.'}, status=status.HTTP_400_BAD_REQUEST)

            valid_versions = [item['value'] for item in version_choices]
            # Endi version tanlash yo‘q: har bir darajada bitta test to‘plami ishlaydi.
            # Eski bazalarda version maydoni qolgan bo‘lsa ham, birinchi mavjud to‘plam olinadi.
            requested_version = valid_versions[0]

        # Agar o‘quvchi testni boshlab, adashib chiqib ketgan bo‘lsa,
        # code qayta kiritilganda bloklamaymiz. Test yakunlanmagan bo‘lsa,
        # aynan o‘sha started_at bilan davom ettiramiz.
        if student.status == Student.Status.IN_PROGRESS:
            started_at = student.started_at
            resume = True
            if not started_at:
                # Eski bazada started_at bo'sh qolgan bo'lsa, 00:00 bo'lib qolmasin.
                started_at = timezone.now()
                resume = False
                student.started_at = started_at
                student.save(update_fields=['started_at'])

            if is_mental_exam:
                return Response(build_mental_payload(started_at, resume=resume))

            payload = build_test_payload(started_at, resume=resume)
            if payload is None:
                return Response({'detail': 'Bu fan va daraja uchun testlar hali qo‘shilmagan.'}, status=status.HTTP_400_BAD_REQUEST)
            return Response(payload)

        # Yangi o‘quvchi code bilan birinchi marta kirganda vaqt aynan hozirdan boshlanadi.
        # Frontend resume=False ni ko‘rib 30:00 dan boshlaydi, 00:00 qilib yubormaydi.
        now = timezone.now()
        student.status = Student.Status.IN_PROGRESS
        student.started_at = now
        student.finished_at = None
        student.is_used = False
        if not is_mental_exam:
            student.selected_version = requested_version or 1
        else:
            student.selected_version = None
        student.progress_answers = {}
        student.progress_remaining_seconds = duration_seconds
        student.progress_current_index = 0
        student.progress_updated_at = timezone.now()
        student.save(update_fields=[
            'status', 'started_at', 'finished_at', 'is_used', 'selected_version', 'progress_answers',
            'progress_remaining_seconds', 'progress_current_index', 'progress_updated_at'
        ])

        if is_mental_exam:
            return Response(build_mental_payload(now, resume=False))

        payload = build_test_payload(now, resume=False)
        if payload is None:
            student.status = Student.Status.NOT_STARTED
            student.started_at = None
            student.selected_version = None
            student.progress_answers = {}
            student.progress_remaining_seconds = None
            student.progress_current_index = 0
            student.progress_updated_at = None
            student.save(update_fields=[
                'status', 'started_at', 'selected_version', 'progress_answers', 'progress_remaining_seconds',
                'progress_current_index', 'progress_updated_at'
            ])
            return Response({'detail': 'Bu fan va daraja uchun testlar hali qo‘shilmagan.'}, status=status.HTTP_400_BAD_REQUEST)

        return Response(payload)


class ExamProgressSaveAPIView(APIView):
    permission_classes = [permissions.AllowAny]

    @transaction.atomic
    def post(self, request):
        code = str(request.data.get('code', '')).strip()
        if not code:
            return Response({'detail': 'Status code yuborilmadi.'}, status=status.HTTP_400_BAD_REQUEST)

        try:
            student = Student.objects.select_related('subject', 'level', 'center').select_for_update().get(code=code)
        except Student.DoesNotExist:
            return Response({'detail': 'Bunday code topilmadi.'}, status=status.HTTP_404_NOT_FOUND)

        if student.status == Student.Status.COMPLETED or student.is_used:
            return Response({'detail': 'Bu test allaqachon yakunlangan.'}, status=status.HTTP_400_BAD_REQUEST)

        if student.status != Student.Status.IN_PROGRESS:
            return Response({'detail': 'Avval testni boshlash kerak.'}, status=status.HTTP_400_BAD_REQUEST)

        is_mental_exam = is_mental_subject(student.subject.name)
        duration_seconds = get_exam_duration_seconds(student)
        answers = normalize_progress_answers(request.data.get('answers', {}), is_mental=is_mental_exam)

        try:
            remaining_seconds = int(request.data.get('remaining_seconds', duration_seconds))
        except (TypeError, ValueError):
            remaining_seconds = get_resume_remaining_seconds(student)
        remaining_seconds = max(0, min(duration_seconds, remaining_seconds))

        try:
            current_index = int(request.data.get('current_index', 0))
        except (TypeError, ValueError):
            current_index = 0
        current_index = max(0, current_index)

        student.progress_answers = answers
        student.progress_remaining_seconds = remaining_seconds
        student.progress_current_index = current_index
        student.progress_updated_at = timezone.now()
        student.save(update_fields=[
            'progress_answers', 'progress_remaining_seconds', 'progress_current_index', 'progress_updated_at'
        ])

        return Response({
            'status': 'ok',
            'remaining_seconds': student.progress_remaining_seconds,
            'saved_answers': student.progress_answers,
            'progress_current_index': student.progress_current_index,
        })


def build_submit_response(result):
    # Yakunlashda qisqa javob qaytariladi: to'liq natija (har bir savol va to'g'ri javoblar)
    # bu yerda kerak emas, uni yig'ish esa yuzlab o'quvchi bir vaqtda yakunlaganda serverni sekinlashtiradi.
    return {
        'status': 'ok',
        'id': result.id,
        'total_questions': result.total_questions,
        'finished_at': result.finished_at,
    }


class ExamSubmitAPIView(APIView):
    permission_classes = [permissions.AllowAny]

    def submit_mental(self, student, answers, remaining_seconds=None):
        task_ids = [item.get('task_id') for item in answers if isinstance(item, dict)]
        tasks = list(MentalTask.objects.select_for_update().filter(student=student, id__in=task_ids).order_by('task_order'))

        # Frontend yuborgan barcha ko‘rsatilgan misollar tekshiriladi.
        # Javob kiritilmagan misol noto‘g‘ri deb saqlanadi.
        answer_map = {}
        for item in answers:
            if not isinstance(item, dict):
                continue
            try:
                task_id = int(item.get('task_id'))
                answer = int(str(item.get('answer', '')).strip())
            except Exception:
                continue
            answer_map[task_id] = answer

        finished_at = timezone.now()
        started_at = student.started_at or finished_at
        spent_seconds = calculate_spent_seconds(student, finished_at, remaining_seconds)

        result = Result.objects.create(
            student=student,
            total_questions=len(tasks),
            correct_count=0,
            percent=0,
            started_at=started_at,
            finished_at=finished_at,
            spent_seconds=spent_seconds,
        )

        correct_count = 0
        for task in tasks:
            student_answer = answer_map.get(task.id)
            is_correct = student_answer == task.correct_answer
            if is_correct:
                correct_count += 1
            task.result = result
            task.student_answer = student_answer
            task.is_correct = is_correct
        MentalTask.objects.bulk_update(tasks, ['result', 'student_answer', 'is_correct'])

        final_score, percent = apply_visible_score_penalty(correct_count, len(tasks))
        result.correct_count = final_score
        result.percent = percent
        result.save(update_fields=['correct_count', 'percent'])

        student.status = Student.Status.COMPLETED
        student.finished_at = finished_at
        student.is_used = True
        student.progress_answers = {}
        student.progress_remaining_seconds = 0
        student.progress_current_index = 0
        student.progress_updated_at = timezone.now()
        student.save(update_fields=[
            'status', 'finished_at', 'is_used', 'progress_answers', 'progress_remaining_seconds',
            'progress_current_index', 'progress_updated_at'
        ])

        return Response(build_submit_response(result), status=status.HTTP_201_CREATED)

    @transaction.atomic
    def post(self, request):
        code = str(request.data.get('code', '')).strip()
        answers = request.data.get('answers', [])

        if not code:
            return Response({'detail': 'Status code yuborilmadi.'}, status=status.HTTP_400_BAD_REQUEST)

        try:
            student = Student.objects.select_related('subject', 'level', 'center').select_for_update().get(code=code)
        except Student.DoesNotExist:
            return Response({'detail': 'Bunday code topilmadi.'}, status=status.HTTP_404_NOT_FOUND)

        if student.status == Student.Status.COMPLETED or student.is_used:
            return Response({'detail': 'Bu code oldin ishlatilgan.'}, status=status.HTTP_400_BAD_REQUEST)

        if student.status != Student.Status.IN_PROGRESS:
            return Response({'detail': 'Avval testni boshlash kerak.'}, status=status.HTTP_400_BAD_REQUEST)

        submitted_remaining_seconds = request.data.get('remaining_seconds', None)
        try:
            submitted_remaining_seconds = int(submitted_remaining_seconds)
        except (TypeError, ValueError):
            submitted_remaining_seconds = None

        if is_mental_subject(student.subject.name):
            return self.submit_mental(student, answers, submitted_remaining_seconds)

        answer_map = {}
        for item in answers:
            try:
                qid = int(item.get('question_id'))
            except Exception:
                continue
            selected = str(item.get('answer', '')).strip().upper()
            if selected in ['A', 'B', 'C', 'D']:
                answer_map[qid] = selected

        exam_version = student.selected_version or 1
        questions = list(Question.objects.filter(subject=student.subject, level=student.level, version=exam_version).order_by('id'))
        correct_count = 0
        finished_at = timezone.now()
        started_at = student.started_at or finished_at
        spent_seconds = calculate_spent_seconds(student, finished_at, submitted_remaining_seconds)

        result = Result.objects.create(
            student=student,
            total_questions=len(questions),
            correct_count=0,
            percent=0,
            started_at=started_at,
            finished_at=finished_at,
            spent_seconds=spent_seconds,
        )

        answer_objects = []
        for question in questions:
            selected = answer_map.get(question.id, '')
            is_correct = selected == question.correct_answer
            if is_correct:
                correct_count += 1
            answer_objects.append(StudentAnswer(
                result=result,
                question=question,
                selected_answer=selected,
                is_correct=is_correct,
            ))
        StudentAnswer.objects.bulk_create(answer_objects)

        final_score, percent = apply_visible_score_penalty(correct_count, len(questions))
        result.correct_count = final_score
        result.percent = percent
        result.save(update_fields=['correct_count', 'percent'])

        student.status = Student.Status.COMPLETED
        student.finished_at = finished_at
        student.is_used = True
        student.progress_answers = {}
        student.progress_remaining_seconds = 0
        student.progress_current_index = 0
        student.progress_updated_at = timezone.now()
        student.save(update_fields=[
            'status', 'finished_at', 'is_used', 'progress_answers', 'progress_remaining_seconds',
            'progress_current_index', 'progress_updated_at'
        ])

        return Response(build_submit_response(result), status=status.HTTP_201_CREATED)
